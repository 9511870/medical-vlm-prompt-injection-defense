"""
Wilcoxon Signed-Rank Analysis — IMAGE-LEVEL VERSION
----------------------------------------------------

Purpose:
    Compare GT-vs-Defense against GT-vs-Attack.

Experimental structure:

    Image
      ├── textual
      │     ├── Run 1
      │     ├── Run 2
      │     └── Run 3
      │
      ├── visual
      │     ├── Run 1
      │     ├── Run 2
      │     └── Run 3
      │
      └── delayed
            ├── Run 1
            ├── Run 2
            └── Run 3


STATISTICAL PROCEDURE
---------------------

Step 1:
    Average the repeated runs for each:

        Image × Injection Type × Condition

Step 2:
    Overall analysis:
        Average textual + visual + delayed
        for each image and condition.

Step 3:
    Run paired Wilcoxon signed-rank test:

        GT vs Defense
              VS
        GT vs Attack

Step 4:
    Also perform separate tests for:

        Textual
        Visual
        Delayed

The statistical unit is the IMAGE.
"""


import numpy as np
import pandas as pd
import openpyxl
import re

from scipy.stats import wilcoxon, rankdata

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter


# ============================================================
# 1. RANK-BISERIAL EFFECT SIZE
# ============================================================

def rank_biserial(a, b):
    """
    Rank-biserial effect size for paired Wilcoxon test.

    a = Defense
    b = Attack

    Range:
        -1 to +1

    Positive:
        Defense tends to be higher than Attack.

    Negative:
        Attack tends to be higher than Defense.

    Zero-difference pairs are removed.
    """

    a = np.asarray(a, dtype=float)
    b = np.asarray(b, dtype=float)

    diff = a - b

    # Remove zero differences
    diff = diff[diff != 0]

    if len(diff) == 0:
        return np.nan

    # Tie-aware ranking of absolute differences
    ranks = rankdata(
        np.abs(diff),
        method="average"
    )

    positive_rank_sum = ranks[diff > 0].sum()
    negative_rank_sum = ranks[diff < 0].sum()

    denominator = (
        positive_rank_sum +
        negative_rank_sum
    )

    if denominator == 0:
        return np.nan

    r = (
        positive_rank_sum -
        negative_rank_sum
    ) / denominator

    return r


# ============================================================
# 2. RUN WILCOXON TEST
# ============================================================

def run_wilcoxon(
    defense,
    attack,
    comparison_name
):

    defense = np.asarray(
        defense,
        dtype=float
    )

    attack = np.asarray(
        attack,
        dtype=float
    )

    # --------------------------------------------------------
    # Remove NaN pairs
    # --------------------------------------------------------

    valid = (
        np.isfinite(defense) &
        np.isfinite(attack)
    )

    defense = defense[valid]
    attack = attack[valid]

    if len(defense) == 0:

        print(
            f"{comparison_name}: "
            f"No paired data."
        )

        return None

    # --------------------------------------------------------
    # Differences
    # --------------------------------------------------------

    difference = defense - attack

    n_total = len(difference)

    n_effective = int(
        np.sum(difference != 0)
    )

    if n_effective == 0:

        print(
            f"{comparison_name}: "
            f"All differences are zero."
        )

        return None

    # --------------------------------------------------------
    # Wilcoxon signed-rank test
    # --------------------------------------------------------

    statistic, p_value = wilcoxon(
        defense,
        attack,
        zero_method="wilcox",
        alternative="two-sided"
    )

    # --------------------------------------------------------
    # Effect size
    # --------------------------------------------------------

    r = rank_biserial(
        defense,
        attack
    )

    # --------------------------------------------------------
    # Median difference
    # --------------------------------------------------------

    median_difference = float(
        np.median(difference)
    )

    # --------------------------------------------------------
    # Effect magnitude
    # --------------------------------------------------------

    abs_r = abs(r)

    if abs_r >= 0.5:

        magnitude = "Large"

    elif abs_r >= 0.3:

        magnitude = "Medium"

    else:

        magnitude = "Small"

    # --------------------------------------------------------
    # Statistical significance
    # --------------------------------------------------------

    if p_value < 0.05:

        significance = (
            "Statistically significant"
        )

    else:

        significance = (
            "Not statistically significant"
        )

    # --------------------------------------------------------
    # Print results
    # --------------------------------------------------------

    print("\n" + "=" * 70)

    print(comparison_name)

    print("=" * 70)

    print(
        f"N (image-level pairs): "
        f"{n_total}"
    )

    print(
        f"Effective N (non-zero differences): "
        f"{n_effective}"
    )

    print(
        f"Median difference "
        f"(Defense - Attack): "
        f"{median_difference:+.4f}"
    )

    print(
        f"Wilcoxon W = "
        f"{statistic:.2f}"
    )

    if p_value < 0.0001:

        print(
            f"p-value = "
            f"{p_value:.4e}"
        )

    else:

        print(
            f"p-value = "
            f"{p_value:.6f}"
        )

    print(
        f"Rank-biserial r = "
        f"{r:+.4f}"
    )

    print(
        f"Effect magnitude = "
        f"{magnitude}"
    )

    print(
        f"Result = "
        f"{significance}"
    )

    return {

        "Comparison":
            comparison_name,

        "N (image-level)":
            int(n_total),

        "N effective":
            int(n_effective),

        "Median Difference (Defense - Attack)":
            median_difference,

        "Wilcoxon W":
            float(statistic),

        "p-value":
            float(p_value),

        "Effect size (rank-biserial r)":
            float(r),

        "Magnitude":
            magnitude,

        "Significance":
            significance
    }


# ============================================================
# 3. LOAD EXCEL FILE
# ============================================================

def load_per_run_scores(filepath):

    wb = openpyxl.load_workbook(
        filepath,
        data_only=True
    )

    if "Per-Run Scores" not in wb.sheetnames:

        raise ValueError(
            "The uploaded Excel file does not contain "
            "'Per-Run Scores' sheet."
        )

    ws = wb["Per-Run Scores"]

    rows = list(
        ws.iter_rows(
            min_row=1,
            values_only=True
        )
    )

    if not rows:

        raise ValueError(
            "The Per-Run Scores sheet is empty."
        )

    # First row = header
    df = pd.DataFrame(
        rows[1:],
        columns=rows[0]
    )

    # --------------------------------------------------------
    # Required columns
    # --------------------------------------------------------

    required_columns = [

        "Original_image_id",
        "Run",
        "Injection Type",
        "Condition",
        "Score"
    ]

    missing = [

        col
        for col in required_columns
        if col not in df.columns
    ]

    if missing:

        raise ValueError(
            f"Missing required columns: {missing}\n\n"
            f"Available columns:\n"
            f"{list(df.columns)}"
        )

    # --------------------------------------------------------
    # Score → numeric
    # --------------------------------------------------------

    df["Score"] = pd.to_numeric(
        df["Score"],
        errors="coerce"
    )

    # --------------------------------------------------------
    # Remove invalid rows
    # --------------------------------------------------------

    df = df.dropna(
        subset=[
            "Original_image_id",
            "Condition",
            "Score"
        ]
    ).copy()

    # --------------------------------------------------------
    # Standardize text
    # --------------------------------------------------------

    df["Condition"] = (
        df["Condition"]
        .astype(str)
        .str.strip()
        .str.lower()
    )

    df["Injection Type"] = (
        df["Injection Type"]
        .astype(str)
        .str.strip()
        .str.lower()
    )

    return df


# ============================================================
# 4. AVERAGE REPEATED RUNS
# ============================================================

def average_runs(df):

    """
    Converts:

        Image × Injection Type × Condition × Run

    into:

        Image × Injection Type × Condition

    by averaging repeated runs.
    """

    grouped = (

        df.groupby(
            [
                "Original_image_id",
                "Injection Type",
                "Condition"
            ],
            as_index=False
        )["Score"]

        .mean()

    )

    return grouped


# ============================================================
# 5. CREATE OVERALL IMAGE-LEVEL SCORES
# ============================================================

def create_overall_image_scores(
    avg_runs
):

    """
    For every image:

        Overall Attack
        Overall Defense

    are calculated by averaging:

        textual
        visual
        delayed

    """

    relevant = avg_runs[
        avg_runs["Condition"].isin(
            [
                "attack",
                "defense"
            ]
        )
    ].copy()

    overall = (

        relevant.groupby(
            [
                "Original_image_id",
                "Condition"
            ],
            as_index=False
        )["Score"]

        .mean()

    )

    pivot = overall.pivot(
        index="Original_image_id",
        columns="Condition",
        values="Score"
    ).reset_index()

    # Keep only complete pairs
    pivot = pivot.dropna(
        subset=[
            "attack",
            "defense"
        ]
    )

    return pivot


# ============================================================
# 6. CREATE IMAGE-LEVEL SCORES BY ATTACK TYPE
# ============================================================

def create_attack_type_scores(
    avg_runs,
    injection_type
):

    """
    Creates:

        Image | Attack | Defense

    for one attack type.
    """

    subset = avg_runs[
        (
            avg_runs["Injection Type"]
            == injection_type
        )
        &
        (
            avg_runs["Condition"].isin(
                [
                    "attack",
                    "defense"
                ]
            )
        )
    ].copy()

    pivot = subset.pivot(
        index="Original_image_id",
        columns="Condition",
        values="Score"
    ).reset_index()

    pivot = pivot.dropna(
        subset=[
            "attack",
            "defense"
        ]
    )

    return pivot


# ============================================================
# 7. INFER MODEL NAME
# ============================================================

def infer_model_name(
    filename
):

    base = filename.rsplit(
        "/",
        1
    )[-1]

    base = re.sub(
        r"\.xlsx$",
        "",
        base,
        flags=re.IGNORECASE
    )

    # Remove trailing (1), (2), _1, etc.
    base = re.sub(
        r"_?\(?\d+\)?$",
        "",
        base
    )

    patterns = [

        r"GT_vs_([A-Za-z0-9]+)_LLM_judge",

        r"^([A-Za-z0-9]+)_LLM_judge",

        r"^([A-Za-z0-9]+)_normal_vs_attack_defense"

    ]

    for pattern in patterns:

        match = re.search(
            pattern,
            base,
            flags=re.IGNORECASE
        )

        if match:

            return match.group(1)

    return "Model"


# ============================================================
# 8. WRITE DATAFRAME USING OPENPYXL
# ============================================================

def write_dataframe_to_sheet(
    ws,
    df
):

    """
    Writes a pandas DataFrame directly into
    an existing openpyxl worksheet.

    This fixes the previous error caused by:

        df.to_excel(ws)

    """

    # --------------------------------------------------------
    # Header
    # --------------------------------------------------------

    for col_idx, column in enumerate(
        df.columns,
        start=1
    ):

        cell = ws.cell(
            row=1,
            column=col_idx,
            value=column
        )

        cell.font = Font(
            bold=True
        )

        cell.fill = PatternFill(
            fill_type="solid",
            fgColor="4472C4"
        )

        cell.font = Font(
            bold=True,
            color="FFFFFF"
        )

        cell.alignment = Alignment(
            horizontal="center"
        )

    # --------------------------------------------------------
    # Data
    # --------------------------------------------------------

    for row_idx, row_data in enumerate(
        df.itertuples(index=False),
        start=2
    ):

        for col_idx, value in enumerate(
            row_data,
            start=1
        ):

            # Convert numpy values to Python values
            if pd.isna(value):

                value = None

            elif isinstance(
                value,
                np.integer
            ):

                value = int(value)

            elif isinstance(
                value,
                np.floating
            ):

                value = float(value)

            ws.cell(
                row=row_idx,
                column=col_idx,
                value=value
            )

    # --------------------------------------------------------
    # Column widths
    # --------------------------------------------------------

    for col_idx, column in enumerate(
        df.columns,
        start=1
    ):

        max_length = len(
            str(column)
        )

        for row_idx in range(
            2,
            len(df) + 2
        ):

            value = ws.cell(
                row=row_idx,
                column=col_idx
            ).value

            if value is not None:

                max_length = max(
                    max_length,
                    len(str(value))
                )

        ws.column_dimensions[
            get_column_letter(col_idx)
        ].width = min(
            max_length + 2,
            35
        )

    ws.freeze_panes = "A2"


# ============================================================
# 9. SAVE ALL RESULTS
# ============================================================

def save_results(
    model_name,
    results,
    overall_scores,
    attack_type_scores,
    output_path
):

    """
    Save all Wilcoxon results and image-level
    intermediate scores into one Excel file.

    Uses openpyxl directly.
    """

    wb = Workbook()

    # ========================================================
    # SHEET 1 — WILCOXON RESULTS
    # ========================================================

    ws = wb.active

    ws.title = "Wilcoxon Results"

    title = (
        f"{model_name} — "
        f"Image-Level Wilcoxon Results"
    )

    ws["A1"] = title

    ws.merge_cells(
        start_row=1,
        start_column=1,
        end_row=1,
        end_column=9
    )

    ws["A1"].font = Font(
        bold=True,
        size=14
    )

    ws["A1"].alignment = Alignment(
        horizontal="center"
    )

    headers = [

        "Comparison",

        "N (Images)",

        "N Effective",

        "Median Difference",

        "W Statistic",

        "p-value",

        "Effect Size (r)",

        "Magnitude",

        "Significance"
    ]

    # Header
    for col_idx, header in enumerate(
        headers,
        start=1
    ):

        cell = ws.cell(
            row=3,
            column=col_idx,
            value=header
        )

        cell.font = Font(
            bold=True,
            color="FFFFFF"
        )

        cell.fill = PatternFill(
            fill_type="solid",
            fgColor="4472C4"
        )

        cell.alignment = Alignment(
            horizontal="center"
        )

    # Results
    row = 4

    for result in results:

        if result is None:
            continue

        values = [

            result[
                "Comparison"
            ],

            result[
                "N (image-level)"
            ],

            result[
                "N effective"
            ],

            result[
                "Median Difference "
                "(Defense - Attack)"
            ],

            result[
                "Wilcoxon W"
            ],

            result[
                "p-value"
            ],

            result[
                "Effect size "
                "(rank-biserial r)"
            ],

            result[
                "Magnitude"
            ],

            result[
                "Significance"
            ]
        ]

        for col_idx, value in enumerate(
            values,
            start=1
        ):

            cell = ws.cell(
                row=row,
                column=col_idx,
                value=value
            )

            # Number formatting
            if col_idx == 4:
                cell.number_format = "0.0000"

            elif col_idx == 5:
                cell.number_format = "0.00"

            elif col_idx == 6:
                cell.number_format = "0.0000E+00"

            elif col_idx == 7:
                cell.number_format = "0.0000"

        row += 1

    # Widths
    widths = [
        42,
        12,
        14,
        20,
        15,
        16,
        20,
        14,
        28
    ]

    for col_idx, width in enumerate(
        widths,
        start=1
    ):

        ws.column_dimensions[
            get_column_letter(col_idx)
        ].width = width

    ws.freeze_panes = "A4"

    # ========================================================
    # SHEET 2 — OVERALL IMAGE SCORES
    # ========================================================

    ws2 = wb.create_sheet(
        "Overall Image Scores"
    )

    write_dataframe_to_sheet(
        ws2,
        overall_scores
    )

    # ========================================================
    # SHEETS 3-5 — ATTACK TYPES
    # ========================================================

    for attack_type, table in (
        attack_type_scores.items()
    ):

        sheet_name = (
            attack_type.capitalize()
        )

        ws_attack = wb.create_sheet(
            sheet_name
        )

        write_dataframe_to_sheet(
            ws_attack,
            table
        )

    # ========================================================
    # SAVE
    # ========================================================

    wb.save(
        output_path
    )

    print("\n")
    print("=" * 70)

    print(
        f"Excel file saved successfully:"
    )

    print(
        output_path
    )

    print("=" * 70)


# ============================================================
# 10. COMPLETE ANALYSIS
# ============================================================

def run_analysis(
    filepath
):

    print("\n")
    print("=" * 70)
    print("LOADING DATA")
    print("=" * 70)

    df = load_per_run_scores(
        filepath
    )

    print(
        f"Raw scored rows: "
        f"{len(df)}"
    )

    print(
        "\nConditions found:"
    )

    print(
        df["Condition"]
        .unique()
        .tolist()
    )

    print(
        "\nInjection types found:"
    )

    print(
        df["Injection Type"]
        .unique()
        .tolist()
    )

    # ========================================================
    # STEP 1 — AVERAGE RUNS
    # ========================================================

    print("\n")
    print("=" * 70)
    print("STEP 1 — AVERAGING REPEATED RUNS")
    print("=" * 70)

    avg_runs = average_runs(
        df
    )

    print(
        f"Records after run averaging: "
        f"{len(avg_runs)}"
    )

    print(
        "\nEach record now represents:"
    )

    print(
        "Image × Injection Type × Condition"
    )

    # ========================================================
    # STEP 2 — OVERALL IMAGE LEVEL
    # ========================================================

    print("\n")
    print("=" * 70)
    print("STEP 2 — OVERALL IMAGE-LEVEL ANALYSIS")
    print("=" * 70)

    overall = (
        create_overall_image_scores(
            avg_runs
        )
    )

    print(
        f"Overall paired images: "
        f"{len(overall)}"
    )

    results = []

    # --------------------------------------------------------
    # Overall Defense vs Attack
    # --------------------------------------------------------

    overall_result = run_wilcoxon(

        defense=overall[
            "defense"
        ].values,

        attack=overall[
            "attack"
        ].values,

        comparison_name=(
            "OVERALL: "
            "GT vs Defense vs "
            "GT vs Attack"
        )
    )

    results.append(
        overall_result
    )

    # ========================================================
    # STEP 3 — INDIVIDUAL ATTACK TYPES
    # ========================================================

    print("\n")
    print("=" * 70)
    print("STEP 3 — ATTACK-TYPE ANALYSIS")
    print("=" * 70)

    attack_types = [

        "textual",

        "visual",

        "delayed"
    ]

    attack_type_tables = {}

    for attack_type in attack_types:

        table = (
            create_attack_type_scores(
                avg_runs,
                attack_type
            )
        )

        attack_type_tables[
            attack_type
        ] = table

        print(
            f"\n{attack_type.upper()} "
            f"paired images: "
            f"{len(table)}"
        )

        if len(table) == 0:

            print(
                f"No complete "
                f"{attack_type} pairs."
            )

            continue

        result = run_wilcoxon(

            defense=table[
                "defense"
            ].values,

            attack=table[
                "attack"
            ].values,

            comparison_name=(
                f"{attack_type.upper()}: "
                f"GT vs Defense vs "
                f"GT vs Attack"
            )
        )

        results.append(
            result
        )

    return (
        results,
        overall,
        attack_type_tables
    )


# ============================================================
# 11. COLAB UPLOAD
# ============================================================

def upload_and_run():

    from google.colab import files

    print(
        "Upload your judge-output Excel file."
    )

    uploaded = files.upload()

    for filename in uploaded:

        print("\n")
        print("=" * 70)

        print(
            f"Uploaded: {filename}"
        )

        print("=" * 70)

        # ----------------------------------------------------
        # Run analysis
        # ----------------------------------------------------

        (
            results,
            overall,
            attack_type_tables
        ) = run_analysis(
            filename
        )

        # ----------------------------------------------------
        # Model name
        # ----------------------------------------------------

        model_name = infer_model_name(
            filename
        )

        print(
            f"\nDetected model: "
            f"{model_name}"
        )

        # ----------------------------------------------------
        # Output filename
        # ----------------------------------------------------

        output_path = (
            f"{model_name}_"
            f"image_level_wilcoxon_results.xlsx"
        )

        # ----------------------------------------------------
        # Save
        # ----------------------------------------------------

        save_results(

            model_name=model_name,

            results=results,

            overall_scores=overall,

            attack_type_scores=
                attack_type_tables,

            output_path=output_path
        )

        # ----------------------------------------------------
        # Download
        # ----------------------------------------------------

        print(
            "\nStarting download..."
        )

        files.download(
            output_path
        )


# ============================================================
# RUN
# ============================================================

if __name__ == "__main__":

    upload_and_run()